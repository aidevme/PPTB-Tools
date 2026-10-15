#!/usr/bin/env python
"""
New-ConsolidatedTemplate.py — generate the blank consolidated canonical matrix.

Emits a blank HSO consolidated requirements-matrix workbook (Legend + Matrix sheets)
from the canonical definition below. This is the **working / review surface** — you
populate it from an RFP, review in one consistent layout, then project back into the
client's own worksheet for submission (it is NOT the client deliverable).

Single source of truth: the canonical lists in this file mirror
references/canonical-model.md. If you change one, change the other.

Usage:
    python New-ConsolidatedTemplate.py [--out consolidated-matrix-template.xlsx]

Exit codes: 0 ok; 2 openpyxl missing.
"""
import sys, argparse

try:
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.stderr.write("openpyxl not installed. Run: python -m pip install -r requirements.txt\n")
    sys.exit(2)

# --- Canonical definition (mirrors references/canonical-model.md) ---
VERDICTS = [
    ("Standard", "Out-of-the-box / native behaviour, no configuration of note."),
    ("Config-basic", "Met via standard user-accessible setup; no code."),
    ("Config-complex", "Met via advanced configuration / simple scripting; no code."),
    ("Modification/Extension", "Requires custom code (plugin, extension, PCF, form script); no core source change."),
    ("Custom-report/query", "Data exists; a custom report/inquiry must be built."),
    ("ISV-embedded", "Point solution embedded in the core app (shared DB/security/UX, no integration build)."),
    ("ISV-pre-integrated", "Separately-marketed product, pre-integrated (separate DB/login, no build)."),
    ("ISV-new", "Point solution that needs an integration to be built."),
    ("Not-supported", "Cannot be met by the proposed solution."),
    ("Future", "On the product roadmap; not available today."),
]
EFFORT_BANDS = [
    ("0", "None — standard / config only"),
    ("1", "Light — minor config-complex or a simple report"),
    ("2", "Moderate — a modest extension or single integration"),
    ("3", "Heavy — significant extension / multi-entity integration"),
    ("4", "Major — net-new build or ISV stand-up + integration"),
]
FIT_STRENGTH = [
    ("high", "Catalog IP directly matches; used in production engagements."),
    ("medium", "Matches but requires customisation or has known limits."),
    ("low", "Adjacent; adapting approaches custom-build cost."),
    ("none", "No catalog IP fits; custom build."),
]
CONFIDENCE = [
    ("auto", "LLM-drafted, not yet reviewed. Never ship at this level."),
    ("reviewed", "A consultant has reviewed the row."),
    ("audited", "Confirmed against product/IP evidence."),
]
DIFFERENTIATOR = ["Yes", "No"]

MATRIX_COLUMNS = [
    "Area", "Sub-area", "Req #", "Requirement", "Client Priority",
    "Canonical Verdict", "Effort Band", "Required Product/ISV", "IP-fit Candidate",
    "Fit Strength", "Differentiator", "Comment", "Confidence/Audited", "Source Ref",
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(italic=True, color="555555")

def _legend_block(ws, start_row, title, rows, widths=(28, 80)):
    ws.cell(start_row, 1, title).font = Font(bold=True, size=12)
    r = start_row + 1
    ws.cell(r, 1, "Value").font = HEADER_FONT
    ws.cell(r, 1).fill = HEADER_FILL
    ws.cell(r, 2, "Meaning").font = HEADER_FONT
    ws.cell(r, 2).fill = HEADER_FILL
    r += 1
    for val, meaning in rows:
        ws.cell(r, 1, val)
        c = ws.cell(r, 2, meaning); c.alignment = Alignment(wrap_text=True)
        r += 1
    return r + 1  # blank line after block

def build(out_path):
    wb = openpyxl.Workbook()

    # --- Legend sheet ---
    legend = wb.active
    legend.title = "Legend"
    legend.cell(1, 1, "HSO Consolidated Requirements Matrix — Legend").font = TITLE_FONT
    legend.cell(2, 1, "Canonical working/review surface. Not the client deliverable — "
                      "project back into the client's own worksheet for submission.").font = SUB_FONT
    legend.column_dimensions["A"].width = 30
    legend.column_dimensions["B"].width = 90
    r = 4
    r = _legend_block(legend, r, "Canonical Verdict", VERDICTS)
    r = _legend_block(legend, r, "Effort Band", EFFORT_BANDS)
    r = _legend_block(legend, r, "Fit Strength", FIT_STRENGTH)
    r = _legend_block(legend, r, "Confidence / Audited", CONFIDENCE)

    # --- Matrix sheet ---
    m = wb.create_sheet("Matrix")
    for idx, name in enumerate(MATRIX_COLUMNS, start=1):
        c = m.cell(1, idx, name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [18, 18, 8, 50, 12, 22, 11, 24, 22, 12, 13, 50, 16, 16]
    for idx, w in enumerate(widths, start=1):
        m.column_dimensions[get_column_letter(idx)].width = w
    m.freeze_panes = "A2"

    # Dropdown validations over a generous row range
    last = 1000
    def add_dv(col_letter, options):
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
        m.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{last}")

    add_dv("F", [v for v, _ in VERDICTS])          # Canonical Verdict
    add_dv("G", [v for v, _ in EFFORT_BANDS])       # Effort Band
    add_dv("J", [v for v, _ in FIT_STRENGTH])       # Fit Strength
    add_dv("K", DIFFERENTIATOR)                      # Differentiator
    add_dv("M", [v for v, _ in CONFIDENCE])         # Confidence/Audited

    wb.save(out_path)
    print(f"Wrote blank consolidated template: {out_path}")
    print("Sheets: Legend, Matrix | dropdowns on Verdict / Effort / Fit Strength / Differentiator / Confidence")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="consolidated-matrix-template.xlsx")
    args = ap.parse_args()
    build(args.out)

if __name__ == "__main__":
    main()
