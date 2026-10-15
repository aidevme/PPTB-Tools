#!/usr/bin/env python
"""
Build an HSO client Word (.docx) deliverable from a JSON content spec, on an HSO base document.

GENERIC — contains no client content. Pass a spec file produced per the project; see
references/word-deliverable.md for the full recipe and gotchas.

Usage:
    python Build-WordDeliverable.py <spec.json>

Spec (all paths absolute):
{
  "base":   "OPTIONAL. Path to an accepted on-brand donor (*_Client.docx / SBD sample) OR the HSO .dotx.
             Omit to use the bundled zero-config HSO donor (assets/HSO_Word_base.docx).
             Format donor only — its body is emptied; never carry the donor's client content.",
  "out":    "output .docx path (do not target a file open in Word)",
  "title":  "<project/report title>",
  "subtitle": "<subtitle>",
  "lead":   "optional one-line intro paragraph",
  "sbd":    false,                   # true => emit the fixed SBD skeleton (front matter + numbered
                                     #         body headings) when no explicit "blocks" are given
  "toc":    true,                    # insert an auto-updating Table of Contents
  "contributors": [                  # HSO attribution — who worked on it. NO single stamped author.
     {"name":"[PLACEHOLDER: name]","role":"[PLACEHOLDER: role]","company":"HSO"}
  ],
  "revision_history": [              # HSO attribution — what changed. NO author column.
     {"version":"1.0","section":"All","change":"Initial Release","date":"<date>"}
  ],
  "blocks": [                        # rendered in order (override the SBD skeleton body)
     {"type":"heading","level":1,"text":"Purpose"},
     {"type":"para","text":"..."},
     {"type":"bullets","items":["...","..."]},
     {"type":"table","headers":["A","B"],"rows":[["1","2"]]}
  ]
  # NOTE: the old "author" + "change_record" (Author = "HSO - <name>" on every row) keys are RETIRED —
  # they wrongly attributed the whole document to one person. Use contributors + revision_history.
}
"""
import sys, json, zipfile, shutil, os, re
from docx import Document

# Bundled zero-config default base: a sanitized on-brand HSO donor (Segoe UI theme, stock styles,
# empty body/header cleaned) derived from the Global HSO Word Template. Used when the spec omits "base".
DEFAULT_DONOR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "HSO_Word_base.docx")
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

# Any [PLACEHOLDER: ...] segment is rendered as a yellow-highlighted run so the reader can see
# at a glance what still needs to be filled in before the deliverable is issued.
PLACEHOLDER_RE = re.compile(r"(\[PLACEHOLDER:[^\]]*\])")

# Non-destructive integrity scan (SOP §6-§7 + references/sbd-qa-checklist.md). Flags authored content that
# looks like (a) fabricated "HSO experience" evidence — the signature the 2026-05 override audit removed — or
# (b) internal-IP / internal-pointer leakage that must not reach a client. Heuristic warnings, NOT a gate.
INTEGRITY_SMELL = [
    (re.compile(r"\b(?:across|in|over)\s+\w+\s+HSO\s+engagements?\b", re.I), "fabrication-smell: counted-engagement claim (SOP §7 — recommendation, not rationale)"),
    (re.compile(r"\b\d+\+?\s+HSO\s+engagements?\b", re.I), "fabrication-smell: counted-engagement claim"),
    (re.compile(r"\bHSO\s+(?:experience|engagements?)\s+(?:show|shows|showed|found|observed)\b", re.I), "fabrication-smell: unsourced 'HSO experience' claim"),
    (re.compile(r"\bHSO\s+Override\b", re.I), "internal-IP leak: 'HSO Override' must not appear in a client document (SOP §7)"),
    (re.compile(r"\bper\s+HSO\s+SOP\b|\bagent/wiki\b|\bADR\s?\d{2,}\b|/SKILL\.md\b", re.I), "internal-pointer leak: SOP/skill/ADR reference (SOP §7)"),
]


def scan_integrity(spec):
    """Walk authored content (lead + blocks) and return [(where, matched_text, message)]."""
    findings = []
    def walk(text, where):
        for rx, msg in INTEGRITY_SMELL:
            for m in rx.finditer(str(text)):
                findings.append((where, m.group(0), msg))
    if spec.get("lead"):
        walk(spec["lead"], "lead")
    for i, b in enumerate(spec.get("blocks", []) or []):
        t = b.get("type")
        if t in ("para", "heading"):
            walk(b.get("text", ""), "blocks[%d] %s" % (i, t))
        elif t == "bullets":
            for j, it in enumerate(b.get("items", []) or []):
                walk(it, "blocks[%d] bullet[%d]" % (i, j))
        elif t == "table":
            for r, row in enumerate(b.get("rows", []) or []):
                for c, cell in enumerate(row):
                    walk(cell, "blocks[%d] table[%d][%d]" % (i, r, c))
    return findings


def add_text(paragraph, text, italic=False):
    """Append text to a paragraph, splitting each [PLACEHOLDER: ...] segment into its own
    yellow-highlighted run. Non-placeholder text renders normally in the paragraph's style.
    Use this instead of add_paragraph(text, ...) / cell.add_run(text) anywhere author-supplied
    content may contain a placeholder (paragraphs, bullets, table cells).

    italic=True renders every run italic -- for asides that are deliberately set apart from the
    body, e.g. an item carried in a section for visibility but excluded from the priced scope."""
    for part in PLACEHOLDER_RE.split(str(text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        if italic:
            run.italic = True
        if PLACEHOLDER_RE.fullmatch(part):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph

# Generalized HSO SBD body skeleton (level, heading). Section 6's per-area sub-headings come from the
# project's in-scope areas via "blocks" — do NOT bake in the full F&O module catalog here.
# Full spec: references/sbd-word-structure.md.
SBD_FRONT_MATTER = ["Copyright", "Trademarks", "Disclaimer"]
SBD_SKELETON = [
    (1, "Introduction"),
    (1, "Solution Blueprint Document Purpose"),
    (1, "Solution Overview"),
    (2, "System Landscape"), (2, "System Integrations/Interfaces"), (2, "Solution Extensions"),
    (1, "System Architecture & Technical Considerations"),
    (2, "Users and Licensing"), (2, "Security"), (2, "Data Management"),
    (2, "Document Management"), (2, "Interface Design Patterns"), (2, "Mobile"),
    (2, "Power Platform apps and services"), (2, "Printing"),
    (1, "Global / Cross-Area Configurations"),
    (1, "Business Solution"),   # + one 2nd-level heading per in-scope area (supply via "blocks")
    (1, "Acceptance"),
]


def open_base(base, work):
    """Open a .docx donor (cloned) or repackage a .dotx (template content-type -> document)."""
    if base.lower().endswith(".dotx"):
        with zipfile.ZipFile(base) as zin, zipfile.ZipFile(work, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename == "[Content_Types].xml":
                    data = data.replace(b"wordprocessingml.template.main+xml",
                                        b"wordprocessingml.document.main+xml")
                zout.writestr(it, data)
    else:
        shutil.copyfile(base, work)
    return Document(work)


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def inject_numbering(doc):
    """Add one decimal multilevel list; return its numId. Headings bind to it (paragraph-level
    numPr overrides any style numbering, so result is single, predictable 1 / 1.1 / 1.1.1)."""
    numbering = doc.part.numbering_part.element
    abs_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    ABS = (max(abs_ids) if abs_ids else 99) + 1
    NUM = (max(num_ids) if num_ids else 99) + 1
    lvls = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        f'<w:lvlText w:val="{t}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="{360*(i+1)}" w:hanging="360"/></w:pPr></w:lvl>'
        for i, t in enumerate(["%1.", "%1.%2", "%1.%2.%3"]))
    abstract = parse_xml(f'<w:abstractNum {W} w:abstractNumId="{ABS}"><w:multiLevelType w:val="multilevel"/>{lvls}</w:abstractNum>')
    existing = numbering.findall(qn("w:abstractNum"))
    (existing[-1].addnext(abstract) if existing else numbering.insert(0, abstract))
    numbering.append(parse_xml(f'<w:num {W} w:numId="{NUM}"><w:abstractNumId w:val="{ABS}"/></w:num>'))
    return NUM


def inject_bullet_numbering(doc):
    """Add one single-level bullet list rendering Word's standard solid black bullet; return numId.

    Needed because the HSO donor ships no "List Bullet" style, and the historic fallback to
    "List Paragraph" is indent-only -- it carries NO bullet glyph, so bulleted content silently
    rendered as plain indented paragraphs. Binding a paragraph-level numPr to a real bullet
    abstractNum makes the glyph appear regardless of which style is used.

    The glyph MUST be F0B7 paired with rFonts="Symbol". Word's solid bullet is Symbol's
    private-use codepoint F0B7, NOT U+2022 or U+25CF -- Symbol has no glyph at those
    codepoints, so pairing either with Symbol renders a hollow missing-glyph box (reported
    from a client review as "open squares", hand-corrected to black bullets in Word). Keep the
    lvlText and the Symbol rFonts together; changing either alone reintroduces the box."""
    numbering = doc.part.numbering_part.element
    abs_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    ABS = (max(abs_ids) if abs_ids else 99) + 1
    NUM = (max(num_ids) if num_ids else 99) + 1
    lvl = ('<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
           '<w:lvlText w:val="&#xF0B7;"/><w:lvlJc w:val="left"/>'
           '<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr>'
           '<w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl>')
    abstract = parse_xml(f'<w:abstractNum {W} w:abstractNumId="{ABS}"><w:multiLevelType w:val="singleLevel"/>{lvl}</w:abstractNum>')
    existing = numbering.findall(qn("w:abstractNum"))
    (existing[-1].addnext(abstract) if existing else numbering.insert(0, abstract))
    numbering.append(parse_xml(f'<w:num {W} w:numId="{NUM}"><w:abstractNumId w:val="{ABS}"/></w:num>'))
    return NUM


def main():
    if len(sys.argv) < 2:
        print("usage: python Build-WordDeliverable.py <spec.json>"); sys.exit(2)
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    out = spec["out"]
    work = os.path.join(os.environ.get("TEMP", "."), "hso_word_base.docx")
    base = spec.get("base") or DEFAULT_DONOR
    if not os.path.exists(base):
        print("ERROR: base not found:", base,
              "\n  (omit 'base' to use the bundled HSO donor, or point it at a donor/.dotx)"); sys.exit(2)
    doc = open_base(base, work)
    styles = set(s.name for s in doc.styles)
    def S(n, fb="Normal"): return n if n in styles else fb
    TITLE = S("Title"); SUB = S("Subtitle", S("AKA Subtitle of Report"))
    HMAP = {1: S("HSO 1st Level Heading", "Heading 1"),
            2: S("HSO 2nd Level Heading", "Heading 2"),
            3: S("HSO 3rd Level Heading", "Heading 3")}
    NORMAL = S("Normal"); BULLET = S("List Bullet", S("List Paragraph"))   # new HSO template lacks "List Bullet"
    TBL = "Table Grid" if "Table Grid" in styles else "Normal Table"

    clear_body(doc)
    try:
        NUM = inject_numbering(doc)
    except Exception as e:
        print("WARN: numbering injection failed; headings will be unnumbered:", e); NUM = None
    # Only needed when we fell back off "List Bullet" -- that style carries its own glyph.
    BNUM = None
    if BULLET != "List Bullet":
        try:
            BNUM = inject_bullet_numbering(doc)
        except Exception as e:
            print("WARN: bullet numbering injection failed; bullets will render unmarked:", e)

    def heading(text, level):
        p = doc.add_paragraph(text, style=HMAP.get(level, HMAP[1]))
        if NUM is not None:
            pPr = p._p.get_or_add_pPr()
            numPr = parse_xml(f'<w:numPr {W}><w:ilvl w:val="{level-1}"/><w:numId w:val="{NUM}"/></w:numPr>')
            ps = pPr.find(qn("w:pStyle")); (ps.addnext(numPr) if ps is not None else pPr.insert(0, numPr))
        return p
    def para(t, italic=False): return add_text(doc.add_paragraph(style=NORMAL), t, italic)
    def bullet(t, italic=False):
        p = add_text(doc.add_paragraph(style=BULLET), t, italic)
        if BNUM is not None:
            pPr = p._p.get_or_add_pPr()
            numPr = parse_xml(f'<w:numPr {W}><w:ilvl w:val="0"/><w:numId w:val="{BNUM}"/></w:numPr>')
            ps = pPr.find(qn("w:pStyle")); (ps.addnext(numPr) if ps is not None else pPr.insert(0, numPr))
        return p
    def bullets(items, italic=False): [bullet(t, italic) for t in items]
    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers)); t.style = TBL
        for i, h in enumerate(headers):
            r = t.rows[0].cells[i].paragraphs[0].add_run(str(h)); r.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row): add_text(cells[i].paragraphs[0], v)
        return t
    def add_toc():
        if "TOC Heading" in styles: doc.add_paragraph("Table of Contents", style="TOC Heading")
        else:
            lp = doc.add_paragraph(style=NORMAL); lp.add_run("Table of Contents").bold = True
        p = doc.add_paragraph()
        for frag in [
            '<w:r %s><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>' % W,
            '<w:r %s><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>' % W,
            '<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % W,
            '<w:r %s><w:t>Right-click and choose "Update Field" (F9) to build the table of contents.</w:t></w:r>' % W,
            '<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % W,
        ]: p._p.append(parse_xml(frag))

    is_sbd = bool(spec.get("sbd"))

    def attribution():
        """Contributors + Revision History — HSO standard. No single stamped author (SOP §5)."""
        contribs = spec.get("contributors")
        if contribs is None and is_sbd:
            contribs = [{"name": "[PLACEHOLDER: name]", "role": "[PLACEHOLDER: role]", "company": "HSO"}]
        if contribs:
            heading("Contributors", 1)
            table(["Name", "Role", "Company"],
                  [[c.get("name", ""), c.get("role", ""), c.get("company", "HSO")] for c in contribs])
        rev = spec.get("revision_history")
        if rev is None and is_sbd:
            rev = [{"version": "1.0", "section": "All", "change": "Initial Release", "date": "<date>"}]
        if rev:
            heading("Revision History", 1)
            rows = [[r.get("version", ""), r.get("section", ""), r.get("change", ""), r.get("date", "")] for r in rev]
            rows += [["", "", "", ""]]
            table(["Version", "Section", "Content Change", "Date"], rows)
        # Retired: "change_record" with an Author column stamped from a single name. Warn if still passed.
        if spec.get("change_record") or spec.get("author"):
            print("WARN: 'change_record'/'author' are retired (single-author stamping). "
                  "Use 'contributors' + 'revision_history'. Ignoring the author stamp.")

    def render_blocks(blocks):
        for b in blocks:
            t = b.get("type")
            if t == "heading": heading(b["text"], int(b.get("level", 1)))
            elif t == "para":  para(b["text"], bool(b.get("italic")))
            elif t == "bullets": bullets(b["items"], bool(b.get("italic")))
            elif t == "table": table(b["headers"], b["rows"])
            elif t == "image":
                if os.path.exists(b.get("path", "")):
                    doc.add_picture(b["path"], width=Inches(b.get("width_in", 6.3)))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if b.get("caption"):
                        cp = doc.add_paragraph(b["caption"], style=NORMAL); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    para(f"[diagram not found: {b.get('path')}]")
            else: print("WARN: unknown block type:", t)

    # --- title block ---
    doc.add_paragraph(spec["title"], style=TITLE)
    if spec.get("subtitle"): doc.add_paragraph(spec["subtitle"], style=SUB)
    if spec.get("lead"): para(spec["lead"])

    # --- SBD front matter (before the TOC): boilerplate placeholders + attribution ---
    if is_sbd:
        for fm in SBD_FRONT_MATTER:
            heading(fm, 1); para(f"[PLACEHOLDER: HSO standard {fm} text]")
    attribution()

    if spec.get("toc"): add_toc(); doc.add_page_break()

    # --- body: explicit blocks override the skeleton; else emit the SBD skeleton (placeholder leaves) ---
    blocks = spec.get("blocks", [])
    if blocks:
        render_blocks(blocks)
    elif is_sbd:
        for level, text in SBD_SKELETON:
            heading(text, level)
            para("[PLACEHOLDER: content for '%s']" % text)

    # --- make Word build/refresh fields (TOC) on open ---
    s = doc.settings.element
    if s.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); s.append(uf)

    # --- clear author metadata: never embed the running maintainer's name (SOP §5) ---
    cp = doc.core_properties
    cp.author = spec.get("core_author", "HSO")
    cp.last_modified_by = spec.get("core_author", "HSO")

    doc.save(out)
    print("SAVED:", out)

    # --- non-destructive integrity review (SOP §6-§7; references/sbd-qa-checklist.md) ---
    smell = scan_integrity(spec)
    if smell:
        print("\nINTEGRITY REVIEW - %d item(s) to check before issuing (non-blocking, heuristic):" % len(smell))
        for where, frag, msg in smell:
            print('  [%s] "%s" - %s' % (where, frag[:60], msg))
        print("  Confirm each is client-appropriate and sourced. Take the recommendation, not internal rationale.")


if __name__ == "__main__":
    main()
